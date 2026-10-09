import os
import re
import unicodedata
import logging
import pdfplumber
from docx import Document

logging.getLogger("pdfminer").setLevel(logging.ERROR)


def extract_text_from_pdf(path):
    with pdfplumber.open(path) as pdf:
        text = ""
        for page in pdf.pages:
            text += (page.extract_text() or "") + "\n"
    return text


def normalize_text(text):
    text = unicodedata.normalize("NFC", text)
    lines = text.split("\n")
    lines = [re.sub(r'\(cid:\d+\)', '', line) for line in lines]
    lines = [re.sub(r'[\U0001F300-\U0001FAFF\u2600-\u27BF]+', '', line) for line in lines]
    lines = [re.sub(r'[ \t\xa0]+', ' ', line).strip() for line in lines]
    lines = [l for l in lines if l]
    return "\n".join(lines)
# def normalize_text(text):
#     text = unicodedata.normalize("NFC", text)
#     lines = text.split("\n")
#     lines = [re.sub(r'\(cid:\d+\)', '', line) for line in lines]
#     lines = [re.sub(r'[\U0001F300-\U0001FAFF\u2600-\u27BF]+', '', line) for line in lines]
#     def _tab_to_colon(line):
#         if ':' in line.split('\t')[0]:
#             return line
#         return line.replace('\t', ': ', 1)

#     lines = [_tab_to_colon(line) for line in lines]
#     lines = [re.sub(r'[ \t\xa0]+', ' ', line).strip() for line in lines]
#     lines = [l for l in lines if l]
#     return "\n".join(lines)

def extract_text_from_docx(path):
    doc = Document(path)
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)

    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2:
                # bảng 2 cột kiểu "Category | Skills" -> chỉ lấy cột giá trị,
                # bỏ cột nhãn để tránh "Category", "Backend"... lọt vào làm skill giả
                table_lines.extend(c for c in cells[1:] if c)
            elif cells and cells[0]:
                table_lines.append(cells[0])
    table_text = "\n".join(table_lines)

    return paragraph_text, table_text


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(path)
    elif ext == ".docx":
        para, table = extract_text_from_docx(path)
        raw_text = para + "\n" + table
    else:
        raise ValueError(f"Không hỗ trợ định dạng: {ext}")
    return normalize_text(raw_text)


def extract_text_sections(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return normalize_text(extract_text_from_pdf(path)), ""
    para, table = extract_text_from_docx(path)
    return normalize_text(para), normalize_text(table)


if __name__ == "__main__":
    text = extract_text("data/nguyễn đức đạt.pdf")
    print(text)