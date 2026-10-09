import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from aggregator import build_dataset, DOMAIN_LABELS

BASE_COLUMNS = [
    "Tệp", "Họ tên", "Số điện thoại", "Email", "LinkedIn", "GitHub",
    "Địa điểm", "Học vấn", "Kinh nghiệm (Tổ chức)",
]
SKILL_COLUMNS = list(DOMAIN_LABELS.values()) + ["Kỹ năng chưa xác định"]
ALL_COLUMNS = BASE_COLUMNS + SKILL_COLUMNS + ["Lỗi"]


def _autofit(ws, max_width=60, min_width=10):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value), default=0)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, min_width), max_width)


def export_to_excel(records, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"

    header = ["STT"] + ALL_COLUMNS
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for i, record in enumerate(records, start=1):
        row = [i] + [record.get(col, "") for col in ALL_COLUMNS]
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    _autofit(ws)

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)


if __name__ == "__main__":
    dataset = build_dataset("data")
    export_to_excel(dataset, "output/ket_qua_trich_xuat_cv.xlsx")
    print(f"Đã xử lý {len(dataset)} CV, xuất báo cáo tại output/ket_qua_trich_xuat_cv.xlsx")
